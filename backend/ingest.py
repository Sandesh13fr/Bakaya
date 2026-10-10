"""Local ingestion: structured records are deterministic; scans fail visibly."""
from __future__ import annotations

import csv
import hashlib
import io
import re
import uuid
from datetime import date
from decimal import Decimal
from pathlib import Path

from .models import Invoice, Payment, SourceRef, Supplier, Buyer, RatePeriod, Message, PurchaseOrder, Delivery, Allocation

MAX_FILE_BYTES = 15 * 1024 * 1024
ALLOWED = {".csv", ".xlsx", ".pdf", ".txt", ".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"}


def parse_date(value):
    # Deliberately do not guess ambiguous dd/mm vs mm/dd dates.
    return date.fromisoformat(str(value).strip()[:10])


def parse_money(value):
    return Decimal(str(value).replace(",", "").replace("₹", "").replace("INR", "").strip())


def row_record(row, sid):
    row = {str(k).strip().lower(): str(v).strip() for k, v in row.items() if k is not None and v is not None and str(v).strip()}
    kind = row.pop("record_type", row.pop("type", ""))
    kwargs = {"source_refs": [sid], "field_sources": {k:[sid] for k in row}}
    if kind == "invoice" or ("invoice_number" in row and "amount" in row):
        if "invoice_number" in row:
            row["number"] = row.pop("invoice_number")
            kwargs["field_sources"]["number"]=[sid]
        if "invoice_date" in row:
            row["date"] = row.pop("invoice_date")
            kwargs["field_sources"]["date"]=[sid]
        row.setdefault("id", row.get("number", f"INV-{uuid.uuid4().hex[:8]}"))
        for key in ("date","delivery_date","acceptance_date","contract_date","objection_date","objection_resolved_date"):
            if key in row:
                row[key]=parse_date(row[key])
        row["amount"]=parse_money(row["amount"])
        if "written_term_days" in row:
            row["written_term_days"]=int(row["written_term_days"])
        return "invoices",Invoice(**kwargs,**row)
    if kind == "payment" or ("date" in row and "amount" in row):
        row.setdefault("id", f"PAY-{uuid.uuid4().hex[:8]}")
        row["date"]=parse_date(row["date"])
        row["amount"]=parse_money(row["amount"])
        for key in ("invoice_ids",):
            if key in row:
                row[key]=[s.strip() for s in row[key].split("|") if s.strip()]
        if "tds_amount" in row:
            row["tds_amount"]=parse_money(row["tds_amount"])
        return "payments",Payment(**kwargs,**row)
    if kind == "supplier":
        if "udyam_date" in row:
            row["udyam_date"]=parse_date(row["udyam_date"])
        return "supplier",Supplier(**kwargs,**row)
    if kind == "buyer":
        return "buyer",Buyer(**kwargs,**row)
    if kind == "rate":
        for key in ("start","end"):
            row[key]=parse_date(row[key])
        row["bank_rate"]=Decimal(row["bank_rate"])
        if "synthetic" in row:
            row["synthetic"]=row["synthetic"].lower() in {"true","1","yes"}
        return "rates",RatePeriod(**kwargs,**row)
    if kind == "delivery":
        row["date"]=parse_date(row["date"])
        if "accepted" in row:
            row["accepted"]=row["accepted"].lower() in {"true","1","yes"}
        return "deliveries",Delivery(**kwargs,**row)
    if kind == "purchase_order":
        row["date"]=parse_date(row["date"])
        return "purchase_orders",PurchaseOrder(**kwargs,**row)
    if kind == "bank_control":
        return "bank_total",parse_money(row["amount"])
    raise ValueError("Unrecognized row. Include record_type and ISO dates; see sample CSV templates.")


def ingest_bytes(filename: str, data: bytes):
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED:
        raise ValueError(f"Unsupported extension {suffix}; accepted: CSV, XLSX, PDF, TXT, images")
    if len(data)>MAX_FILE_BYTES:
        raise ValueError("File exceeds15 MB limit")
    root = "SRC-"+uuid.uuid4().hex[:12]
    base = SourceRef(id=root,file=Path(filename).name,kind="uploaded_document",sha256=hashlib.sha256(data).hexdigest(),stored_name=root+suffix)
    sources, extracted, errors = [base], {}, []
    if suffix in {".png",".jpg",".jpeg",".webp",".tif",".tiff"}:
        base.status="ocr_unavailable"
        base.text="OCR is not configured. Re-upload text-searchable evidence or provide a structured CSV with verified facts. No date or amount inferred."
        return sources,extracted,errors
    if suffix in {".csv",".xlsx"}:
        if suffix == ".csv":
            text=data.decode("utf-8-sig")
            table=list(csv.DictReader(io.StringIO(text)))
        else:
            from openpyxl import load_workbook
            workbook=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
            sheet=workbook.active
            rows=list(sheet.iter_rows(values_only=True))
            if not rows:
                raise ValueError("Workbook is empty")
            headers=[str(v or "").strip() for v in rows[0]]
            table=[dict(zip(headers,row)) for row in rows[1:]]
            workbook.close()
            text="\n".join(",".join(str(v or "") for v in row.values()) for row in table)
        base.text=text[:100000]
        base.kind="structured_ledger"
        for n,row in enumerate(table,2):
            sid=f"{root}-L{n}"
            source=base.model_copy(update={"id":sid,"line":n,"text":str(row),"stored_name":base.stored_name})
            sources.append(source)
            try:
                key,record=row_record(row,sid)
                if key in {"supplier","buyer","bank_total"}:
                    extracted[key]=record
                else:
                    extracted.setdefault(key,[]).append(record)
            except Exception as exc:
                source.status="parse_error"
                errors.append(f"Row{n}: {exc}")
        return sources,extracted,errors
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data))
        pages=[]
        for n,page in enumerate(reader.pages,1):
            text=page.extract_text() or ""
            sid=f"{root}-P{n}"
            source=base.model_copy(update={"id":sid,"page":n,"text":text[:100000],"stored_name":base.stored_name})
            source.status="unstructured_review" if text.strip() else "ocr_unavailable"
            sources.append(source)
            pages.append(text)
        base.text="\n".join(pages)[:100000]
        base.status="parsed" if base.text.strip() else "ocr_unavailable"
        return sources,extracted,errors
    text=data.decode("utf-8-sig")
    base.text=text[:100000]
    base.kind="message_export"
    # TXT is evidence only: message text cannot prove delivery/acceptance automatically.
    extracted["messages"]=[]
    for n,line in enumerate(text.splitlines(),1):
        if not line.strip():
            continue
        sid=f"{root}-L{n}"
        sources.append(base.model_copy(update={"id":sid,"line":n,"text":line}))
        extracted["messages"].append(Message(id=f"MSG-{root}-{n}",text=line,source_refs=[sid],field_sources={"text":[sid]}))
    return sources,extracted,errors
